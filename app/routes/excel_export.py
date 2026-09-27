import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ดึงฟังก์ชันรายงานยอดขายประจำวัน
try:
    from app.database.report_db import get_today_report
except ImportError:
    get_today_report = None


def create_excel(event_name: str = "") -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Report"

    # 1. กำหนด Header
    headers = ["ลำดับ", "วัน-เวลา", "หมวดหมู่/รายการ", "ไซส์", "ราคา (บาท)", "ช่องทางการชำระ", "จุดขาย/งาน"]
    ws.append(headers)

    # 2. ดึงข้อมูลจาก Database
    data = []
    if get_today_report:
        try:
            data = get_today_report(event_name) or []
        except Exception as e:
            print(f"Error fetching sales report data: {e}")

    # 3. วนลูปนำข้อมูลใส่ Sheet
    for idx, item in enumerate(data, 1):
        if isinstance(item, dict):
            row = [
                idx,
                item.get("created_at") or item.get("date") or "-",
                item.get("category") or item.get("title") or "-",
                item.get("size") or "-",
                float(item.get("price") or item.get("amount") or 0.0),
                item.get("payment_method") or "เงินสด",
                item.get("station_name") or item.get("event_name") or event_name or "-"
            ]
        elif isinstance(item, (list, tuple)):
            row = [
                idx,
                item[1] if len(item) > 1 else "-",
                item[2] if len(item) > 2 else "-",
                item[3] if len(item) > 3 else "-",
                float(item[4]) if len(item) > 4 else 0.0,
                item[5] if len(item) > 5 else "เงินสด",
                item[6] if len(item) > 6 else (event_name or "-")
            ]
        else:
            row = [idx, "-", "-", "-", 0.0, "-", event_name or "-"]

        ws.append(row)

    # ----------------------------------------------------
    # 🎨 4. จัดสไตล์ให้สวยงาม + ล็อกแถวหัวข้อ
    # ----------------------------------------------------
    # ❄️ ล็อกแถวที่ 1 ไม่ให้เลื่อนตามเมื่อ Scroll
    ws.freeze_panes = "A2"
    ws.views.sheetView[0].showGridLines = True

    # นิยาม Font, สี และเส้นขอบ
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")  # สีกรมท่าเข้ม
    row_fill_even = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")  # สีสลับแถว

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    # ตกแต่ง Header Row (แถวที่ 1)
    ws.row_dimensions[1].height = 28
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border

    # ตกแต่ง Data Rows (แถวที่ 2 เป็นต้นไป)
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 22
        is_even = (row_idx % 2 == 0)

        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if is_even:
                cell.fill = row_fill_even

            # จัด Format คอลัมน์
            if col_idx == 5:  # คอลัมน์ราคา
                cell.alignment = align_right
                cell.number_format = "#,##0.00"
            elif col_idx in [1, 2, 4, 6]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # ปรับความกว้างคอลัมน์อัตโนมัติ (คำนวณภาษาไทย)
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            length = sum(2 if ord(char) > 127 else 1 for char in val_str)
            max_len = max(max_len, length)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 12)

    # 5. บันทึกไฟล์
    filepath = "temp_sales_report.xlsx"
    wb.save(filepath)
    return filepath